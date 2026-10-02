"""Regression tests for the public exporter CLI and generated workbook."""
import csv
import json
import os
from pathlib import Path
import runpy
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook

SCRIPT = Path(os.environ.get('IDEA_SKILL_ROOT', Path(__file__).resolve().parents[1] / 'idea-evaluation')) / 'scripts/build_xlsx.py'
CRIT = ('upside', 'demand', 'feasibility', 'cost', 'speed', 'fit')


def idea(**extra):
    return dict(idea='Example', **dict(zip(CRIT, [4, 5, 5, 5, 1, 1])), **extra)


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def export(self, data, csv_mode=False, expect_ok=True):
        src = self.root / 'input.json'
        out = self.root / ('out.csv' if csv_mode else 'out.xlsx')
        src.write_text(json.dumps(data), encoding='utf-8')
        result = subprocess.run([sys.executable, str(SCRIPT), str(src), str(out)] + (['--csv'] if csv_mode else []), capture_output=True, text=True)
        self.assertEqual(result.returncode == 0, expect_ok, result.stderr)
        return out

    def rows(self, data):
        with self.export(data, True).open() as f:
            return list(csv.DictReader(f))

    def test_unrounded_priority_and_workbook_formulas(self):
        data = {'ideas': [idea(evidence='E2')]}
        self.assertEqual(self.rows(data)[0]['priority'], 'B')  # 3.375 displayed 3.4
        src = self.root / 'input.json'; src.write_text(json.dumps(data))
        out = self.root / 'formula.xlsx'
        with patch.object(sys, 'argv', [str(SCRIPT), str(src), str(out)]):
            namespace = runpy.run_path(str(SCRIPT))
        self.assertEqual(namespace['score'](data['ideas'][0]), (3.375, 'B'))
        wb = load_workbook(out); ws = wb['Comparison']
        self.assertNotIn('ROUND(', ws['N2'].value)
        self.assertNotIn('ROUND(', ws['O2'].value)
        self.assertIn('O2>=Settings!$B$18', ws['P2'].value)
        self.assertEqual(ws['O2'].number_format, '0.0')
        self.assertEqual(len(ws.conditional_formatting), 1)

    def test_exact_threshold_and_custom_settings(self):
        d = dict(idea='Threshold', evidence='E0', **dict.fromkeys(CRIT, 4))
        rows = self.rows({'thresholds': {'A': 3.2}, 'ideas': [d]})
        self.assertEqual(rows[0]['priority'], 'A')
        wb = load_workbook(self.export({'thresholds': {'A': 3.2}, 'ideas': [d]}))
        self.assertEqual(wb['Settings']['B18'].value, 3.2)

    def test_stopped_incomplete_and_default_evidence(self):
        rows = self.rows({'ideas': [idea(gate_feasibility='no'), {'idea': 'Missing'}, idea()]})
        self.assertEqual([r['priority'] for r in rows], ['Stopped', 'incomplete', 'B'])
        self.assertTrue(all(r['evidence'] == 'E0' for r in rows))
        wb = load_workbook(self.export({'ideas': [idea()]}))
        self.assertEqual(wb['Comparison']['M2'].value, 'E0')

    def test_formula_like_text_is_literal(self):
        for prefix in ('=', '+', '-', '@', ' \t='):
            with self.subTest(prefix=prefix):
                d = idea(problem=prefix+'HYPERLINK("https://example.org","x")')
                self.assertTrue(self.rows({'ideas': [d]})[0]['problem'].startswith("'"))
                wb = load_workbook(self.export({'ideas': [d], 'notes': ['=1+1']}))
                self.assertEqual(wb['Comparison']['C2'].data_type, 's')
                self.assertEqual(wb['Comparison']['C2'].value, d['problem'])
                cells = [c for row in wb['Comparison'] for c in row if c.value == '=1+1']
                self.assertEqual(cells[0].data_type, 's')

    def test_decimal_break_even_and_nonpositive_margin(self):
        rows = self.rows({'ideas': [idea(price=.3, variable_cost=.1, fixed_costs=.6), idea(price=1, variable_cost=2, fixed_costs=10)]})
        self.assertEqual(rows[0]['breakeven_customers'], '3')
        self.assertEqual(rows[1]['breakeven_customers'], '')

    def test_reject_invalid_input_without_output(self):
        bad = [[], {'ideas': []}, {'ideas': [{'idea': ''}]},
               {'ideas': [idea(evidence='E99')]}, {'ideas': [idea(gate_feasibility='maybe')]},
               {'ideas': [idea(hours_week=-1)]}, {'ideas': [dict(idea(), upside=True)]},
               {'ideas': [dict(idea(), upside=6)]}, {'ideas': [dict(idea(), upside=2.5)]},
               {'ideas': [dict(idea(), upside=float('nan'))]},
               {'ideas': [idea()], 'weights': {'surprise': .1}},
               {'ideas': [idea()], 'weights': {'upside': .8}},
               {'ideas': [idea()], 'evidence_factors': {'E2': .2}},
               {'ideas': [idea()], 'thresholds': {'B': 4}},
               {'ideas': [idea()], 'time_budget_h_week': -1},
               {'ideas': [idea()], 'notes': '=1+1'}]
        for data in bad:
            with self.subTest(data=data):
                for mode in (False, True):
                    out = self.root / ('out.csv' if mode else 'out.xlsx')
                    out.unlink(missing_ok=True)
                    self.export(data, mode, False)
                    self.assertFalse(out.exists())

    def test_research_goal_and_dynamic_capacity(self):
        wb = load_workbook(self.export({'goal': 'research', 'time_budget_h_week': 6, 'ideas': [idea(hours_week=3)]}))
        self.assertEqual(wb['Settings']['B24'].value, 'research')
        self.assertEqual(wb['Capacity']['B1'].value, '=Settings!B22')
        self.assertIn('SUMIF', wb['Capacity']['B2'].value)
        self.assertIn('ABS(SUM(B2:B7)-1)', wb['Settings']['B25'].value)


if __name__ == '__main__':
    unittest.main()
