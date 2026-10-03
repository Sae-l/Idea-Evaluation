"""Unit tests for scoring.py and build_xlsx.py. Run: python tests/test_scoring.py
Needs openpyxl for the spreadsheet part (pip install -r requirements-dev.txt); use --skip-xlsx to skip it explicitly."""
import csv, itertools, json, os, subprocess, sys, tempfile
HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(HERE, "..", "skills", "idea-evaluation", "scripts")
sys.path.insert(0, SCRIPTS)
from scoring import CRIT, EVIDENCE, InputError, THRESHOLDS, WEIGHTS, config, economics, normalize, score, sensitivity, validate

data = json.load(open(os.path.join(HERE, "sample_ideas.json")))
w, ev, th = config(data)
by = {i["idea"]: i for i in data["ideas"]}
idea = lambda **k: {"idea": "x", **dict(zip(CRIT, (3,) * 6)), **k}

# --- hand-computed cases
assert score(by["Course"], w, ev, th) == (3.5, "A")          # raw 3.85 * 0.90 = 3.465 -> 3.5
assert score(by["Newsletter"], w, ev, th) == (3.2, "B")      # 4.0 * 0.8
assert score(by["Device"], w, ev, th) == (2.2, "C")
assert score(by["Marketplace"], w, ev, th) == (2.6, "C")
assert score(by["Blocked"], w, ev, th) == (None, "Stopped")  # a gate 'no' overrides everything
assert score(by["Partial"], w, ev, th) == (None, "incomplete")
assert score(by["Floor"], w, ev, th) == (4.2, "C")           # would be A, demand = 1 caps it
assert economics(by["Course"]) == (37, 17)                   # 600 / 37 = 16.2 -> 17
assert economics(by["Newsletter"]) == (None, None)
assert economics({"price": 20, "variable_cost": 5}) == (15, None)         # margin without fixed costs
assert economics({"price": 10, "variable_cost": 12, "fixed_costs": 100}) == (-2, None)
assert abs(sum(w.values()) - 1) < 1e-9

# --- half-up rounding like the spreadsheet (binary round() gives 0.8 here)
assert score(idea(**dict(zip(CRIT, (1,) * 6)), evidence="E1"))[0] == 0.9

# --- thresholds exactly at the boundary (E4 = factor 1.0, no rating of 1 so the floor is not involved)
seen = {}
for r in itertools.product(range(2, 6), repeat=6):
    s, p = score({"idea": "x", **dict(zip(CRIT, r)), "evidence": "E4"})
    seen.setdefault(s, p)
for value, expected in ((3.4, "A"), (3.3, "B"), (2.8, "B"), (2.7, "C"), (2.2, "C"), (2.1, "D")):
    assert seen.get(value) == expected, (value, seen.get(value), expected)

# --- knockout floor on each protected criterion; not on the others
for crit in ("demand", "feasibility", "cost"):
    assert score(idea(**{crit: 1, **{c: 5 for c in CRIT if c != crit}, "evidence": "E4"}))[1] == "C", crit
assert score(idea(**{"upside": 1, **{c: 5 for c in CRIT if c != "upside"}, "evidence": "E4"}))[1] == "A"

# --- case-insensitive gates and evidence
assert score(idea(gate_viability="No")) == (None, "Stopped")
assert score(idea(evidence=" e4 ")) == score(idea(evidence="E4"))
assert normalize({"gate_viability": " NO "})["gate_viability"] == "no"

# --- validation
bad = {"ideas": [idea(upside="4"), idea(demand=9), idea(idea="", fit=3), {"upside": 3}, idea(gate_viability="maybe"),
                 idea(evidence="E9"), idea(price="cheap"), idea(), idea()], "weights": {"upside": .5}, "time_budget_h_week": -1}
for k in (0, 1, 4, 5, 6):
    bad["ideas"][k]["idea"] = f"case{k}"          # unique names; items 7 and 8 stay duplicates on purpose
errs = validate(bad)
joined = "\n".join(errs)
for needle in ("'upside' must be a whole number from 1 to 5", "'demand' must be a whole number from 1 to 5", "'idea' (name) is required",
               "gate_viability", "'evidence'", "'price' must be a finite number", "duplicate name", "weights must sum to 1", "time_budget_h_week"):
    assert needle in joined, (needle, errs)
assert validate({"ideas": [idea(upside=True)]}) != []          # booleans are not numbers
assert validate({"ideas": [idea(upside=3.5)]}) != []           # ratings are whole numbers (the sheet only accepts those)
for bad_extra in ({"price": float("nan")}, {"price": float("inf")}, {"price": -1}, {"fixed_costs": -5}, {"hours_week": -2},
                  {"notes": {"x": 1}}, {"problem": ["a"]}):
    assert validate({"ideas": [idea(**bad_extra)]}), bad_extra
assert validate({"ideas": [idea(), idea(idea="X ")], "time_budget_h_week": float("inf")}) != []
assert validate({"ideas": [idea(idea="a"), idea(idea="A ")]}) != []                        # duplicates ignore case and spaces
assert validate({"ideas": [idea()], "currency": 5}) != []
assert validate({"ideas": "nope"}) and validate([]) and validate({"ideas": []}) == []
try:
    config({"weights": {"nonsense": 1}}); raise SystemExit("config accepted unknown key")
except InputError:
    pass
try:
    config({"thresholds": {"A": 2, "B": 3}}); raise SystemExit("config accepted unordered thresholds")
except InputError:
    pass
for bad_cfg in ({"weights": {"upside": float("nan")}}, {"evidence_factors": {"E0": float("inf")}}, {"weights": {"upside": -.1, "fit": .35}}):
    try:
        config(bad_cfg); raise SystemExit(f"config accepted {bad_cfg}")
    except InputError:
        pass

# --- custom configuration is honoured
w2, ev2, th2 = config({"weights": {"upside": .35, "fit": 0}, "thresholds": {"A": 4.0}})
assert abs(sum(w2.values()) - 1) < 1e-9 and th2["A"] == 4.0 and th2["B"] == THRESHOLDS["B"]

# --- sensitivity: 12 weight runs + 2 evidence + 2 threshold runs; duplicate names are fine; ties are handled
four = [by[k] for k in ("Course", "Newsletter", "Device", "Marketplace")]
s = sensitivity(four, w, ev, th)
assert s["runs"] == 16 and s["top"] == ["Course"], s
assert sensitivity([idea(), idea()], w, ev, th)["top"] == ["x", "x"]      # exact tie keeps both
assert sensitivity([], w, ev, th)["runs"] == 16
# the "top" idea follows the export order: priority class first (a knockout idea is capped at C), then score
X = idea(idea="X knockout", upside=5, demand=1, feasibility=5, cost=5, speed=5, fit=5, evidence="E4")
Y = idea(idea="Y solid", upside=4, demand=4, feasibility=4, cost=4, speed=4, fit=4, evidence="E4")
assert score(X, w, ev, th) == (4.2, "C") and score(Y, w, ev, th) == (4.0, "A")
assert sensitivity([X, Y], w, ev, th)["top"] == ["Y solid"], sensitivity([X, Y], w, ev, th)
print("scoring ok;", s)

# --- command line and spreadsheet export
skip = "--skip-xlsx" in sys.argv
try:
    import openpyxl
except ImportError:
    openpyxl = None
    if not skip:
        sys.exit("openpyxl missing: run pip install -r requirements-dev.txt (or pass --skip-xlsx to skip explicitly)")


def run(*args, expect=0):
    r = subprocess.run([sys.executable, *args], capture_output=True, text=True)
    assert (r.returncode == 0) == (expect == 0), (args, r.returncode, r.stdout, r.stderr)
    return r


with tempfile.TemporaryDirectory() as d:
    p = lambda name: os.path.join(d, name)
    sample = os.path.join(HERE, "sample_ideas.json")
    # scoring CLI: ok, invalid input -> exit 2 with a message naming the row, flag before file name
    run(os.path.join(SCRIPTS, "scoring.py"), "--sensitivity", sample)
    json.dump({"ideas": [idea(upside="4")]}, open(p("bad.json"), "w"))
    r = run(os.path.join(SCRIPTS, "scoring.py"), p("bad.json"), expect=2)
    assert r.returncode == 2 and "idea #1" in r.stderr and "Traceback" not in r.stderr, r.stderr
    open(p("broken.json"), "w").write("{not json")
    open(p("nan.json"), "w").write('{"ideas": [{"idea": "x", "price": NaN}]}')
    assert "not allowed" in run(os.path.join(SCRIPTS, "scoring.py"), p("nan.json"), expect=2).stderr
    assert "cannot read" in run(os.path.join(SCRIPTS, "scoring.py"), p("broken.json"), expect=2).stderr

    # CSV export: all columns, sorted like the sheet, BOM, injection neutralized, input never overwritten
    inj = {"currency": "=CUR", "ideas": [idea(idea="=HYPERLINK(\"http://x\")", problem="@SUM(1)", notes="-cmd", gate_viability="yes",
                                               price=10, variable_cost=12, fixed_costs=100, cost_to_mvp="low"),
                                          idea(idea="b", upside=5, demand=5, feasibility=5, cost=5, speed=5, fit=5, evidence="E4")]}
    json.dump(inj, open(p("inj.json"), "w"))
    run(os.path.join(SCRIPTS, "build_xlsx.py"), p("inj.json"), p("i.csv"), "--csv")
    raw = open(p("i.csv"), "rb").read()
    assert raw.startswith(b"\xef\xbb\xbf"), "CSV needs a UTF-8 BOM"
    rows = list(csv.DictReader(open(p("i.csv"), encoding="utf-8-sig")))
    assert [r["idea"] for r in rows][0] == "b", "rows must be sorted by priority/score"
    other = rows[1]
    assert other["idea"].startswith("'=") and other["problem"].startswith("'@") and other["notes"].startswith("'-"), other
    assert other["gate_viability"] == "yes" and other["breakeven_customers"] == "no margin", other
    assert other["cost_to_mvp"] == "low", other   # accepted input field must not be dropped from the export
    run(os.path.join(SCRIPTS, "build_xlsx.py"), p("inj.json"), p("inj.json"), "--csv", expect=2)      # same in/out path
    run(os.path.join(SCRIPTS, "build_xlsx.py"), p("bad.json"), p("x.csv"), "--csv", expect=2)          # invalid input
    json.dump({"ideas": []}, open(p("empty.json"), "w"))
    run(os.path.join(SCRIPTS, "build_xlsx.py"), p("empty.json"), p("e.csv"), "--csv")

    if openpyxl:
        inj.update({"time_budget_h_week": 6, "notes": ["=1+1", "@note"]})
        json.dump(inj, open(p("inj2.json"), "w"))
        run(os.path.join(SCRIPTS, "build_xlsx.py"), p("inj2.json"), p("i.xlsx"))
        run(os.path.join(SCRIPTS, "build_xlsx.py"), p("empty.json"), p("e.xlsx"))                       # empty list must not crash
        wb = openpyxl.load_workbook(p("i.xlsx"))
        assert wb.sheetnames == ["Comparison", "Settings", "Capacity"]
        assert wb["Settings"]["B2"].value == 0.25 and wb["Settings"]["B18"].value == 3.4 and wb["Settings"]["B22"].value == 6
        user_cells = [c for sh in wb for row in sh.iter_rows() for c in row
                      if isinstance(c.value, str) and c.value.startswith(("=HYPERLINK", "@SUM", "=1+1", "@note", "-cmd", "=CUR"))]
        assert len(user_cells) >= 5 and all(c.data_type == "s" for c in user_cells), [(c.coordinate, c.data_type) for c in user_cells]
        assert any(c.data_type == "s" and "CUR" in str(c.value) for c in wb["Comparison"][1]), "currency header must stay text"
        print("build ok; injection neutralized in idea, problem, notes, currency and header cells")
