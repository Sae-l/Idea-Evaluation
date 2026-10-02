"""Build a comparison workbook (or CSV) from JSON. Currency- and user-neutral.

Usage: python build_xlsx.py ideas.json output.xlsx [--csv]
  --csv   write output as CSV (no dependencies) instead of xlsx.

ideas.json schema (all fields optional except "idea"):
{
  "time_budget_h_week": 8,
  "currency": "EUR",
  "goal": "profit",
  "weights": {"upside": .25, "demand": .20, "feasibility": .15, "cost": .15, "speed": .15, "fit": .10},
  "evidence_factors": {"E0": .80, "E1": .85, "E2": .90, "E3": .95, "E4": 1.0},
  "thresholds": {"A": 3.4, "B": 2.8, "C": 2.2},
  "ideas": [{
    "id": 1, "idea": "...", "problem": "who has which problem; today solved by ...",
    "gate_desirability": "yes|no|unknown", "gate_feasibility": "...", "gate_viability": "...",
    "upside": 4, "demand": 3, "feasibility": 4, "cost": 5, "speed": 4, "fit": 3,
    "evidence": "E0", "hours_week": 5, "weeks_to_first_evidence": 2,
    "cost_to_mvp": "low", "riskiest_assumption": "...", "test": "...", "pass_threshold": "...",
    "price": 20, "variable_cost": 5, "fixed_costs": 600, "premortem": "...", "notes": ""
  }],
  "notes": ["Estimates are assumptions; state their basis and uncertainty."]
}
Scores are recomputed in the sheet from weights, evidence factors and thresholds (yellow cells are editable).
"""
import csv, json, sys
import math
from decimal import Decimal


def numeric(value, label, lo=0, hi=None):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"{label} must be a finite number")
    if value < lo or (hi is not None and value > hi):
        raise ValueError(f"{label} is outside its allowed range")


def checked_update(defaults, supplied, label):
    if not isinstance(supplied, dict) or set(supplied) - set(defaults):
        raise ValueError(f"{label}: unknown keys or invalid object")
    return dict(defaults, **supplied)


def safe_csv(value):
    if isinstance(value, str) and value.lstrip(' \t\r\n').startswith(('=', '+', '-', '@')):
        return "'" + value
    return value

args = [a for a in sys.argv[1:] if not a.startswith("--")]
as_csv = "--csv" in sys.argv
if len(args) != 2:
    sys.exit(__doc__)
with open(args[0], encoding="utf-8") as source:
    data = json.load(source)
if not isinstance(data, dict):
    sys.exit("Input must be a JSON object")
out = args[1]

W = {"upside": .25, "demand": .20, "feasibility": .15, "cost": .15, "speed": .15, "fit": .10}
W = checked_update(W, data.get("weights", {}), "weights")
EF = {"E0": .80, "E1": .85, "E2": .90, "E3": .95, "E4": 1.0}
EF = checked_update(EF, data.get("evidence_factors", {}), "evidence_factors")
TH = {"A": 3.4, "B": 2.8, "C": 2.2}
TH = checked_update(TH, data.get("thresholds", {}), "thresholds")
CRIT = list(W)
ideas = data.get("ideas", [])
for key, value in W.items():
    numeric(value, key, 0, 1)
if abs(sum(W.values()) - 1) > 1e-9:
    raise ValueError("weights must sum to 1")
for key, value in EF.items():
    numeric(value, key, 0, 1)
if list(EF.values()) != sorted(EF.values()):
    raise ValueError("evidence factors must be nondecreasing")
for key, value in TH.items():
    numeric(value, key, 0, 5)
if not TH["A"] > TH["B"] > TH["C"]:
    raise ValueError("thresholds must satisfy A > B > C")
numeric(data.get("time_budget_h_week", 8), "time budget")
if data.get("goal", "profit") not in ("profit", "impact", "research", "learning", "other"):
    raise ValueError("invalid goal")
if not isinstance(data.get("currency", ""), str):
    raise ValueError("currency must be text")
if not isinstance(data.get("notes", []), list) or any(not isinstance(n, str) for n in data.get("notes", [])):
    raise ValueError("notes must be a list of strings")
if not isinstance(ideas, list) or not ideas:
    raise ValueError("ideas must contain at least one idea")
for i, d in enumerate(ideas):
    if not isinstance(d, dict) or not isinstance(d.get("idea"), str) or not d["idea"].strip():
        raise ValueError(f"idea {i + 1}: nonempty idea text required")
    for field in ("problem", "cost_to_mvp", "riskiest_assumption", "test", "pass_threshold", "premortem", "notes"):
        if d.get(field) is not None and not isinstance(d[field], str):
            raise ValueError(f"{field} must be text")
    if d.get("id") is not None and (isinstance(d["id"], bool) or not isinstance(d["id"], (str, int))):
        raise ValueError("id must be text or an integer")
    for c in CRIT:
        if d.get(c) is not None:
            numeric(d[c], c, 1, 5)
            if int(d[c]) != d[c]:
                raise ValueError(f"{c} must be a whole-number rating")
    for g in ("gate_desirability", "gate_feasibility", "gate_viability"):
        if d.get(g) not in (None, "yes", "no", "unknown"):
            raise ValueError(f"invalid {g}")
    if d.get("evidence", "E0") not in EF:
        raise ValueError("evidence must be E0 through E4")
    for field in ("hours_week", "weeks_to_first_evidence", "price", "variable_cost", "fixed_costs"):
        if d.get(field) is not None:
            numeric(d[field], field)


def score(d):
    if any(d.get(g) == "no" for g in ("gate_desirability", "gate_feasibility", "gate_viability")):
        return None, "Stopped"
    vals = [d.get(c) for c in CRIT]
    if any(v is None for v in vals):
        return None, "incomplete"
    raw = sum(Decimal(str(W[c])) * Decimal(str(d[c])) for c in CRIT)
    adj = raw * Decimal(str(EF[d.get("evidence", "E0")]))
    p = next((key for key in ("A", "B", "C") if adj >= Decimal(str(TH[key]))), "D")
    return float(adj), p


def econ(d):
    try:
        m = Decimal(str(d["price"])) - Decimal(str(d["variable_cost"]))
        be = math.ceil(Decimal(str(d["fixed_costs"])) / m) if m > 0 else None
        return float(m), be
    except (KeyError, TypeError, ArithmeticError):
        return None, None


if as_csv:
    cols = ["id", "idea", "problem"] + CRIT + ["evidence", "adjusted_score", "priority", "hours_week",
            "weeks_to_first_evidence", "riskiest_assumption", "test", "pass_threshold", "margin", "breakeven_customers"]
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(cols)
        for d in ideas:
            s, p = score(d); m, be = econ(d)
            row = [d.get("id"), d.get("idea"), d.get("problem")] + [d.get(c) for c in CRIT] + \
                  [d.get("evidence", "E0"), round(s, 1) if s is not None else None, p,
                   d.get("hours_week"), d.get("weeks_to_first_evidence"), d.get("riskiest_assumption"),
                   d.get("test"), d.get("pass_threshold"), m, be]
            w.writerow([safe_csv(v) for v in row])
    print("saved:", out); sys.exit()

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

F = Font(name="Arial", size=10); HF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F4E78"); IN = PatternFill("solid", fgColor="FFF2CC")
S = Side(style="thin", color="BFBFBF"); BD = Border(top=S, bottom=S, left=S, right=S)
PRIO = {"A": "C6EFCE", "B": "FFEB9C", "C": "FCE4D6", "D": "F2F2F2"}
cur = data.get("currency", "")

wb = Workbook()
# Settings sheet first so formulas can reference it
st = wb.active; st.title = "Settings"
rows = [("Criterion", "Weight")] + [(c.capitalize(), W[c]) for c in CRIT] + [("Sum (must be 100 %)", "=SUM(B2:B7)"),
        (None, None), ("Evidence level", "Factor")] + [(k, EF[k]) for k in ("E0", "E1", "E2", "E3", "E4")] + \
       [(None, None), ("Priority threshold (adjusted score from)", None), ("A", TH["A"]), ("B", TH["B"]), ("C", TH["C"]),
        (None, None), ("Time budget (hours/week)", data.get("time_budget_h_week", 8))]
for r in rows: st.append(r)
for row in st.iter_rows():
    for c in row: c.font = F
for r in (1, 10):
    for c in st[r]: c.font = HF; c.fill = HFILL
for r in list(range(2, 8)) + list(range(11, 16)) + [18, 19, 20, 22]: st[f"B{r}"].fill = IN
for r in range(2, 9): st[f"B{r}"].number_format = "0%"
st.column_dimensions["A"].width = 42; st.column_dimensions["B"].width = 12
st["A24"] = "Goal"; st["B24"] = data.get("goal", "profit")
st["A25"] = "Scoring settings valid"
st["B25"] = '=AND(COUNT(B2:B7)=6,ABS(SUM(B2:B7)-1)<0.000000001,MIN(B2:B7)>=0,MAX(B2:B7)<=1,COUNT(B11:B15)=5,MIN(B11:B15)>=0,MAX(B11:B15)<=1,B11<=B12,B12<=B13,B13<=B14,B14<=B15,COUNT(B18:B20)=3,B18>B19,B19>B20,B20>=0,B18<=5)'
st["A27"] = "Scores are heuristic preferences, not success probabilities."
st["A28"] = "Interpret demand and evidence relative to the goal."


cols = [("ID", "id", 5), ("Idea", "idea", 30), ("Problem / who / today solved by", "problem", 34),
        ("Gate: desirability", "gate_desirability", 11), ("Gate: feasibility", "gate_feasibility", 11),
        ("Gate: viability", "gate_viability", 11)] + [(c.capitalize() + " (1-5)", c, 9) for c in CRIT] + \
       [("Evidence (E0-E4)", "evidence", 9), ("Raw score", "=raw", 8), ("Adjusted score", "=adj", 9),
        ("Priority", "=prio", 8), ("Hours/week", "hours_week", 8), ("Weeks to first evidence", "weeks_to_first_evidence", 10),
        ("Riskiest assumption", "riskiest_assumption", 32), ("Cheapest test", "test", 36), ("Pass threshold", "pass_threshold", 24),
        (f"Price {cur}".strip(), "price", 8), (f"Variable cost {cur}".strip(), "variable_cost", 10),
        (f"Fixed costs {cur}".strip(), "fixed_costs", 10), ("Margin", "=margin", 8), ("Break-even customers", "=be", 10),
        ("Pre-mortem", "premortem", 34), ("Notes", "notes", 28)]
K = {k: L(i) for i, (_, k, _) in enumerate(cols, 1)}
ws = wb.create_sheet("Comparison", 0)
ws.append([c[0] for c in cols])
for c in ws[1]: c.font = HF; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical="center")
for i, (_, _, w) in enumerate(cols, 1): ws.column_dimensions[L(i)].width = w


def fx(key, r):
    c = lambda k: f"{K[k]}{r}"
    gates = [c("gate_desirability"), c("gate_feasibility"), c("gate_viability")]
    if key == "=raw":
        refs = [c(x) for x in CRIT]
        w = "+".join(f"{x}*Settings!$B${i}" for i, x in enumerate(refs, 2))
        return f'=IF(COUNT({refs[0]}:{refs[-1]})<{len(CRIT)},"",{w})'
    if key == "=adj":
        return (f'=IF(OR({c("=raw")}="",NOT(Settings!$B$25)),"",IF(OR({gates[0]}="no",{gates[1]}="no",{gates[2]}="no"),"",'
                f'{c("=raw")}*IFERROR(VLOOKUP({c("evidence")},Settings!$A$11:$B$15,2,FALSE),Settings!$B$11)))')
    if key == "=prio":
        return (f'=IF(OR({gates[0]}="no",{gates[1]}="no",{gates[2]}="no"),"Stopped",IF({c("=adj")}="","incomplete",'
                f'IF({c("=adj")}>=Settings!$B$18,"A",IF({c("=adj")}>=Settings!$B$19,"B",IF({c("=adj")}>=Settings!$B$20,"C","D")))))')
    if key == "=margin":
        return f'=IF(COUNT({c("price")},{c("variable_cost")})=2,{c("price")}-{c("variable_cost")},"")'
    if key == "=be":
        return f'=IF(AND(ISNUMBER({c("=margin")}),ISNUMBER({c("fixed_costs")})),IF({c("=margin")}>0,ROUNDUP({c("fixed_costs")}/{c("=margin")},0),"no margin"),"")'


order = {"A": 0, "B": 1, "C": 2, "D": 3}
ideas_sorted = sorted(ideas, key=lambda d: (order.get(score(d)[1], 4), -(score(d)[0] or 0)))
for r, d in enumerate(ideas_sorted, 2):
    for j, (_, k, _) in enumerate(cols, 1):
        cell = ws.cell(r, j, fx(k, r) if k.startswith("=") else d.get(k, "E0" if k == "evidence" else None))
        if not k.startswith("=") and isinstance(cell.value, str):
            cell.data_type = "s"
        if k in ("=raw", "=adj"):
            cell.number_format = "0.0"
        cell.font = F; cell.border = BD; cell.alignment = Alignment(wrap_text=True, vertical="top")
        if k in CRIT or k in ("evidence", "hours_week", "price", "variable_cost", "fixed_costs") or k.startswith("gate_"):
            cell.fill = IN

last = len(ideas_sorted) + 1
for p, color in PRIO.items():
    ws.conditional_formatting.add(f"{K['=prio']}2:{K['=prio']}{last}", CellIsRule(operator="equal", formula=[f'"{p}"'], fill=PatternFill("solid", fgColor=color)))
dv = DataValidation(type="list", formula1='"yes,no,unknown"', allow_blank=True); ws.add_data_validation(dv)
dv.showErrorMessage = True; dv.errorStyle = "stop"
dv.add(f"{K['gate_desirability']}2:{K['gate_viability']}{last}")
dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True); ws.add_data_validation(dv2)
dv2.showErrorMessage = True; dv2.errorStyle = "stop"
dv2.add(f"{K[CRIT[0]]}2:{K[CRIT[-1]]}{last}")
dv3 = DataValidation(type="list", formula1='"E0,E1,E2,E3,E4"', allow_blank=True); ws.add_data_validation(dv3)
dv3.showErrorMessage = True; dv3.errorStyle = "stop"
dv3.add(f"{K['evidence']}2:{K['evidence']}{last}")
ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:{L(len(cols))}{last}"
for k, t in enumerate(["Notes:", "Yellow cells are inputs; scores and priorities recalculate from the Settings sheet."] + data.get("notes", [])):
    cell = ws.cell(last + 2 + k, 2, t)
    cell.data_type = "s"
    cell.font = Font(name="Arial", size=10, bold=(k == 0), italic=(k > 0))

# Capacity
cp = wb.create_sheet("Capacity")
rng = lambda col: f"Comparison!${col}$2:${col}${last}"
pc, hw = K["=prio"], K["hours_week"]
for row in [("Available hours per week", "=Settings!B22"), ("Hours/week of all A ideas", f'=SUMIF({rng(pc)},"A",{rng(hw)})'),
            ("Hours/week of all B ideas", f'=SUMIF({rng(pc)},"B",{rng(hw)})'), ("Sum A + B", "=B2+B3"),
            ("Load", '=IF(B1>0,B4/B1,"")'),
            ("Assessment", '=IF(B1="","",IF(B2>B1,"A ideas alone exceed capacity: run one at a time",IF(B4>B1,"Overloaded: park B ideas","fits")))')]:
    cp.append(row)
for row in cp.iter_rows():
    for c in row: c.font = F
cp["B5"].number_format = "0%"; cp.column_dimensions["A"].width = 34; cp.column_dimensions["B"].width = 48

wb.save(out)
print("saved:", out)
