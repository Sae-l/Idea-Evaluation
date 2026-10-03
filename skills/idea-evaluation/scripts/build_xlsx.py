"""Build a comparison workbook (or CSV) from JSON. Currency- and user-neutral.
Input is validated first (whole ratings 1-5, finite non-negative numbers, known gates and evidence levels, unique names); invalid input exits with code 2
and a message naming the problem. Text cells are never interpreted as formulas.

Usage: python build_xlsx.py ideas.json output.xlsx [--csv]
  --csv   write output as CSV (no dependencies) instead of xlsx.

ideas.json schema (all fields optional except "idea"):
{
  "time_budget_h_week": 8,
  "currency": "EUR",
  "weights": {"upside": .25, "demand": .20, "feasibility": .15, "cost": .15, "speed": .15, "fit": .10},
  "evidence_factors": {"E0": .80, "E1": .85, "E2": .90, "E3": .95, "E4": 1.0},
  "thresholds": {"A": 3.4, "B": 2.8, "C": 2.2},
  "ideas": [{
    "id": 1, "idea": "...", "problem": "who has which problem; today solved by ...",
    "gate_desirability": "yes|no|unknown", "gate_feasibility": "...", "gate_viability": "...",
    "upside": 4, "demand": 3, "feasibility": 4, "cost": 5, "speed": 4, "fit": 3,
    "evidence": "E0", "hours_week": 5, "weeks_to_first_evidence": 2,
    "cost_to_mvp": "low", "riskiest_assumption": "...", "test": "...", "pass_threshold": "...",
    "price": 20, "variable_cost": 5, "fixed_costs": 600, "premortem": "...", "notes": ""
  }],
  "notes": ["All values are rough estimates (+/-50 %)."]
}
Scores are recomputed in the sheet from weights, evidence factors and thresholds (yellow cells are editable).
"""
import csv, os, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from scoring import CRIT, breakeven_label, config, economics as econ, load, normalize, score as _score, sort_key, validate

args = [a for a in sys.argv[1:] if not a.startswith("--")]
as_csv = "--csv" in sys.argv
if len(args) != 2:
    print(__doc__, file=sys.stderr); sys.exit(2)
if os.path.abspath(args[0]) == os.path.abspath(args[1]):
    print("output path must differ from the input path", file=sys.stderr); sys.exit(2)
data = load(args[0])
errs = validate(data)
if errs:
    print("invalid input:\n  " + "\n  ".join(errs), file=sys.stderr); sys.exit(2)
out = args[1]

W, EF, TH = config(data)
ideas = [normalize(d) for d in data["ideas"]]
score = lambda d: _score(d, W, EF, TH)
ideas_sorted = sorted(ideas, key=lambda d: sort_key(d, W, EF, TH))

DANGEROUS = ("=", "+", "-", "@", "\t", "\r")


def safe_csv(v):
    """Neutralize spreadsheet formula injection in text cells (OWASP: prefix with a single quote)."""
    return "'" + v if isinstance(v, str) and v.startswith(DANGEROUS) else v


if as_csv:
    cols = (["id", "idea", "problem", "gate_desirability", "gate_feasibility", "gate_viability"] + CRIT +
            ["evidence", "adjusted_score", "priority", "cost_to_mvp", "hours_week", "weeks_to_first_evidence", "riskiest_assumption",
             "test", "pass_threshold", "price", "variable_cost", "fixed_costs", "margin", "breakeven_customers",
             "premortem", "notes"])
    with open(out, "w", newline="", encoding="utf-8-sig") as f:   # BOM so spreadsheet software reads UTF-8
        w = csv.writer(f); w.writerow(cols)
        for d in ideas_sorted:
            s_, p = score(d); m, _ = econ(d)
            vals = {**d, "adjusted_score": s_, "priority": p, "margin": m, "breakeven_customers": breakeven_label(d)}
            w.writerow([safe_csv(vals.get(c)) for c in cols])
    print("saved:", out); sys.exit()

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

F = Font(name="Arial", size=10); HF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F4E78"); IN = PatternFill("solid", fgColor="FFF2CC")
S = Side(style="thin", color="BFBFBF"); BD = Border(top=S, bottom=S, left=S, right=S)
PRIO = {"A": "C6EFCE", "B": "FFEB9C", "C": "FCE4D6", "D": "F2F2F2"}
cur = str(data.get("currency", ""))


def put(sheet, r, c, v):
    """Write a user-supplied value; text is never interpreted as a formula."""
    cell = sheet.cell(r, c, v)
    if isinstance(v, str) and v.startswith(DANGEROUS):
        cell.data_type = "s"
    return cell

wb = Workbook()
# Settings sheet first so formulas can reference it
st = wb.active; st.title = "Settings"
rows = [("Criterion", "Weight")] + [(c.capitalize(), W[c]) for c in CRIT] + [("Sum (must be 100 %)", "=SUM(B2:B7)"),
        (None, None), ("Evidence level", "Factor")] + [(k, EF[k]) for k in ("E0", "E1", "E2", "E3", "E4")] + \
       [(None, None), ("Priority threshold (adjusted score from)", None), ("A", TH["A"]), ("B", TH["B"]), ("C", TH["C"]),
        (None, None), ("Time budget (hours/week)", float(data.get("time_budget_h_week", 8))),
        ("Usable share of the time budget", 0.7)]
for r in rows: st.append(r)
for row in st.iter_rows():
    for c in row: c.font = F
for r in (1, 10):
    for c in st[r]: c.font = HF; c.fill = HFILL
for r in list(range(2, 8)) + list(range(11, 16)) + [18, 19, 20, 22, 23]: st[f"B{r}"].fill = IN
for r in list(range(2, 9)) + [23]: st[f"B{r}"].number_format = "0%"
st.column_dimensions["A"].width = 42; st.column_dimensions["B"].width = 12
GATE_ROWS = None

cols = [("ID", "id", 5), ("Idea", "idea", 30), ("Problem / who / today solved by", "problem", 34),
        ("Gate: desirability", "gate_desirability", 11), ("Gate: feasibility", "gate_feasibility", 11),
        ("Gate: viability", "gate_viability", 11)] + [(c.capitalize() + " (1-5)", c, 9) for c in CRIT] + \
       [("Evidence (E0-E4)", "evidence", 9), ("Raw score", "=raw", 8), ("Adjusted score", "=adj", 9),
        ("Priority", "=prio", 8), ("Cost to MVP", "cost_to_mvp", 10), ("Hours/week", "hours_week", 8), ("Weeks to first evidence", "weeks_to_first_evidence", 10),
        ("Riskiest assumption", "riskiest_assumption", 32), ("Cheapest test", "test", 36), ("Pass threshold", "pass_threshold", 24),
        (f"Price {cur}".strip(), "price", 8), (f"Variable cost {cur}".strip(), "variable_cost", 10),
        (f"Fixed costs {cur}".strip(), "fixed_costs", 10), ("Margin", "=margin", 8), ("Break-even customers", "=be", 10),
        ("Pre-mortem", "premortem", 34), ("Notes", "notes", 28)]
K = {k: L(i) for i, (_, k, _) in enumerate(cols, 1)}
ws = wb.create_sheet("Comparison", 0)
for j, c_ in enumerate(cols, 1): put(ws, 1, j, c_[0])
for c in ws[1]: c.font = HF; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical="center")
for i, (_, _, w) in enumerate(cols, 1): ws.column_dimensions[L(i)].width = w


def fx(key, r):
    c = lambda k: f"{K[k]}{r}"
    gates = [c("gate_desirability"), c("gate_feasibility"), c("gate_viability")]
    if key == "=raw":
        refs = [c(x) for x in CRIT]
        w = "+".join(f"{x}*Settings!$B${i}" for i, x in enumerate(refs, 2))
        return f'=IF(COUNT({refs[0]}:{refs[-1]})<{len(CRIT)},"",ROUND({w},2))'
    if key == "=adj":
        return (f'=IF({c("=raw")}="","",IF(OR({gates[0]}="no",{gates[1]}="no",{gates[2]}="no"),"",'
                f'ROUND({c("=raw")}*IFERROR(VLOOKUP({c("evidence")},Settings!$A$11:$B$15,2,FALSE),Settings!$B$11),1)))')
    if key == "=prio":
        letter = (f'IF({c("=adj")}>=Settings!$B$18,"A",IF({c("=adj")}>=Settings!$B$19,"B",IF({c("=adj")}>=Settings!$B$20,"C","D")))')
        floor = f'MIN({c("demand")},{c("feasibility")},{c("cost")})<=1'
        return (f'=IF(OR({gates[0]}="no",{gates[1]}="no",{gates[2]}="no"),"Stopped",IF({c("=adj")}="","incomplete",'
                f'IF(AND({floor},OR({letter}="A",{letter}="B")),"C",{letter})))')
    if key == "=margin":
        return f'=IF(COUNT({c("price")},{c("variable_cost")})=2,{c("price")}-{c("variable_cost")},"")'
    if key == "=be":
        return f'=IF(AND(ISNUMBER({c("=margin")}),ISNUMBER({c("fixed_costs")})),IF({c("=margin")}>0,ROUNDUP({c("fixed_costs")}/{c("=margin")},0),"no margin"),"")'


for r, d in enumerate(ideas_sorted, 2):
    for j, (_, k, _) in enumerate(cols, 1):
        cell = put(ws, r, j, fx(k, r)) if k.startswith("=") else put(ws, r, j, d.get(k))
        if k.startswith("="):
            cell.data_type = "f"
        cell.font = F; cell.border = BD; cell.alignment = Alignment(wrap_text=True, vertical="top")
        if k in CRIT or k in ("evidence", "hours_week", "price", "variable_cost", "fixed_costs") or k.startswith("gate_"):
            cell.fill = IN
    p = score(d)[1]
    if p in PRIO: ws[f"{K['=prio']}{r}"].fill = PatternFill("solid", fgColor=PRIO[p])
last = max(len(ideas_sorted) + 1, 2)
dv = DataValidation(type="list", formula1='"yes,no,unknown"', allow_blank=True); ws.add_data_validation(dv)
dv.add(f"{K['gate_desirability']}2:{K['gate_viability']}{last}")
dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True); ws.add_data_validation(dv2)
dv2.add(f"{K[CRIT[0]]}2:{K[CRIT[-1]]}{last}")
dv3 = DataValidation(type="list", formula1='"E0,E1,E2,E3,E4"', allow_blank=True); ws.add_data_validation(dv3)
dv3.add(f"{K['evidence']}2:{K['evidence']}{last}")
ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:{L(len(cols))}{last}"
for k, t in enumerate(["Notes:", "Yellow cells are inputs; scores and priorities recalculate from the Settings sheet. Knockout floor: a rating of 1 for demand, feasibility or cost caps the priority at C."] + data.get("notes", [])):
    put(ws, last + 2 + k, 2, t).font = Font(name="Arial", size=10, bold=(k == 0), italic=(k > 0))

# Capacity
cp = wb.create_sheet("Capacity")
rng = lambda col: f"Comparison!${col}$2:${col}${last}"
pc, hw = K["=prio"], K["hours_week"]
for row in [("Available hours per week", "=Settings!B22"), ("Usable hours (share set in Settings)", "=B1*Settings!B23"),
            ("Hours/week of all A ideas", f'=SUMIF({rng(pc)},"A",{rng(hw)})'),
            ("Hours/week of all B ideas", f'=SUMIF({rng(pc)},"B",{rng(hw)})'), ("Sum A + B", "=B3+B4"),
            ("Load (of usable hours)", '=IF(B2>0,B5/B2,"")'),
            ("Assessment", '=IF(B2="","",IF(B3>B2,"A ideas alone exceed the usable hours: run one at a time",IF(B5>B2,"Overloaded: park B ideas","fits")))')]:
    cp.append(row)
for row in cp.iter_rows():
    for c in row: c.font = F
cp["B6"].number_format = "0%"; cp.column_dimensions["A"].width = 34; cp.column_dimensions["B"].width = 48

wb.save(out)
print("saved:", out)
